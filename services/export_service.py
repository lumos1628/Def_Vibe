import csv
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import Config
from repositories.chat_wheel_repository import ChatWheelRepository
from repositories.hero_repository import HeroRepository
from repositories.match_repository import MatchRepository
from services.drive_service import DriveService


class ExportService:
    CHAT_FIELDNAMES = [
        "match_id",
        "chat_time",
        "chat_time_min",
        "chat_message",
        "chat_type",
        "chat_is_chatwheel",
        "chat_wheel_id",
        "sender_hero_id",
        "sender_hero_name",
        "sender_slot",
        "sender_player_slot",
        "has_replay_data",
        "unparseable",
    ]

    KILL_FIELDNAMES = [
        "match_id",
        "kill_time",
        "kill_time_min",
        "killer_hero_id",
        "killer_hero_name",
        "victim_hero_id",
        "victim_hero_name",
        "rampage_time",
        "rampage_time_min",
        "is_rampage_kill",
        "has_replay_data",
        "unparseable",
    ]

    def __init__(
        self,
        hero_repository: HeroRepository,
        match_repository: MatchRepository,
        chat_wheel_repository: ChatWheelRepository,
        drive_service: DriveService | None = None,
    ):
        self.hero_repo = hero_repository
        self.match_repo = match_repository
        self.chat_wheel_repo = chat_wheel_repository
        self.drive_service = drive_service
        heroes = self.hero_repo.list_all()
        self.hero_names = {h["id"]: h["display_name"] for h in heroes}
        self.hero_short_to_id = {
            h["short_name"]: h["id"]
            for h in heroes
            if h.get("short_name")
        }

    def _hero_name(self, hero_id: int | None) -> str:
        if hero_id is None:
            return ""
        return self.hero_names.get(hero_id, "")

    def _victim_hero_id(self, internal_name: str) -> int | None:
        if internal_name.startswith("npc_dota_hero_"):
            short_name = internal_name[len("npc_dota_hero_"):]
        else:
            short_name = internal_name
        return self.hero_short_to_id.get(short_name)

    def _resolve_chat_message(self, msg: dict) -> tuple[str, bool, int | None]:
        msg_type = msg.get("type", "")
        key = msg.get("key", "")
        if msg_type != "chatwheel":
            return str(key), False, None

        try:
            chat_wheel_id = int(key)
        except (ValueError, TypeError):
            return str(key), True, None

        entry = self.chat_wheel_repo.get_by_id(chat_wheel_id)
        if entry is None:
            return f"chatwheel:{key}", True, chat_wheel_id

        message = entry.get("label") or entry.get("message") or entry.get("name") or f"chatwheel:{key}"
        return message, True, chat_wheel_id

    def _detect_rampage(self, match_id: int, data: dict) -> tuple[int | None, list[int]]:
        hero_ids = self.match_repo.get_hero_ids_for_match(match_id)
        if not hero_ids:
            return None, []

        players = {p.get("hero_id"): p for p in data.get("players") or []}
        for hero_id in hero_ids:
            player = players.get(hero_id)
            if not player:
                continue

            kills_log = sorted(
                player.get("kills_log") or [],
                key=lambda k: k.get("time", 0),
            )
            if len(kills_log) < 5:
                continue

            for i in range(len(kills_log) - 4):
                window = kills_log[i : i + 5]
                times = [k.get("time", 0) for k in window]
                if times[-1] - times[0] <= 18:
                    return times[0], times
        return None, []

    def _player_slot_to_hero_id(self, data: dict) -> dict[int, int]:
        return {
            p["player_slot"]: p["hero_id"]
            for p in data.get("players") or []
            if p.get("player_slot") is not None and p.get("hero_id") is not None
        }

    def export_chat_kills(
        self, output_path: Path | None = None, today: str | None = None
    ) -> tuple[Path, Path, Path]:
        today = today or datetime.now(timezone.utc).date().isoformat()
        export_dir = output_path or Config.EXPORT_DIR / today
        export_dir.mkdir(parents=True, exist_ok=True)

        chats_csv_path = export_dir / "chats.csv"
        kills_csv_path = export_dir / "kills.csv"
        xlsx_path = export_dir / f"rampage_export_{today}.xlsx"

        if not chats_csv_path.exists() and not kills_csv_path.exists():
            cleared = self.match_repo.clear_export_log_for_date(today)
            if cleared:
                print(f"\nExport del día no encontrado. Se limpiaron {cleared} registro(s) de exportación de {today}.")

        match_details = self.match_repo.list_matches_to_export(today)
        if not match_details:
            print(f"\nNo hay partidas encontradas y parseadas hoy ({today}) pendientes de exportar.")
            return chats_csv_path, kills_csv_path, xlsx_path

        print(f"\nExportando {len(match_details)} partida(s) del día {today}...")

        chat_rows: list[dict] = []
        kill_rows: list[dict] = []

        for match_id, data, has_replay_data, unparseable in match_details:
            chat_rows.extend(
                self._build_chat_rows(match_id, data, has_replay_data, unparseable)
            )
            kill_rows.extend(
                self._build_kill_rows(match_id, data, has_replay_data, unparseable)
            )

        if not chat_rows and not kill_rows:
            print("No se generaron filas de exportación.")
            return chats_csv_path, kills_csv_path, xlsx_path

        self._write_csv(chats_csv_path, self.CHAT_FIELDNAMES, chat_rows)
        self._write_csv(kills_csv_path, self.KILL_FIELDNAMES, kill_rows)
        match_ids = [match_id for match_id, _, _, _ in match_details]
        self._write_excel(xlsx_path, chat_rows, kill_rows, match_ids)

        now = datetime.now(timezone.utc).isoformat()
        for match_id, _, _, _ in match_details:
            self.match_repo.mark_exported(match_id, now, str(export_dir))

        print(f"\n✓ Exportadas {len(chat_rows)} filas de chats a {chats_csv_path}")
        print(f"✓ Exportadas {len(kill_rows)} filas de kills a {kills_csv_path}")
        print(f"✓ Exportado Excel a {xlsx_path}")
        return chats_csv_path, kills_csv_path, xlsx_path

    def export_chat_kills_csv(self, output_path: Path | None = None) -> Path:
        _, _, xlsx_path = self.export_chat_kills(output_path)
        return xlsx_path

    def export_match_ids_csv(
        self, match_ids: list[int], today: str | None = None
    ) -> Path:
        """Exports a CSV with one match ID per row, named with today's date."""
        today = today or datetime.now(timezone.utc).date().isoformat()
        csv_path = Config.EXPORT_DIR / f"match_ids_{today}.csv"
        csv_path.parent.mkdir(parents=True, exist_ok=True)

        fieldnames = ["match_id"]
        rows = [{"match_id": match_id} for match_id in match_ids]
        self._write_csv(csv_path, fieldnames, rows)

        print(f"\n✓ Exportadas {len(rows)} IDs a {csv_path}")
        if self.drive_service is not None:
            self.drive_service.upload(csv_path)
        return csv_path

    def reexport_parsed_today(self, output_path: Path | None = None) -> tuple[Path, Path, Path]:
        today = datetime.now(timezone.utc).date().isoformat()
        cleared = self.match_repo.clear_export_log_for_parsed_today(today)
        print(f"\nSe limpiaron {cleared} registro(s) de exportación de partidas parseadas hoy ({today}).")
        return self.export_chat_kills(output_path, today)

    def _build_chat_rows(
        self, match_id: int, data: dict, has_replay_data: bool, unparseable: bool
    ) -> list[dict]:
        player_slots = self._player_slot_to_hero_id(data)
        rows: list[dict] = []

        for msg in data.get("chat") or []:
            time_value = msg.get("time")
            if time_value is None:
                continue

            message, is_chatwheel, chat_wheel_id = self._resolve_chat_message(msg)
            player_slot = msg.get("player_slot")
            sender_hero_id = (
                player_slots.get(player_slot)
                if player_slot is not None
                else None
            )

            rows.append(
                {
                    "match_id": match_id,
                    "chat_time": time_value,
                    "chat_time_min": round(time_value / 60, 2),
                    "chat_message": message,
                    "chat_type": msg.get("type", ""),
                    "chat_is_chatwheel": 1 if is_chatwheel else 0,
                    "chat_wheel_id": chat_wheel_id,
                    "sender_hero_id": sender_hero_id,
                    "sender_hero_name": self._hero_name(sender_hero_id),
                    "sender_slot": msg.get("slot"),
                    "sender_player_slot": player_slot,
                    "has_replay_data": 1 if has_replay_data else 0,
                    "unparseable": 1 if unparseable else 0,
                }
            )

        return rows

    def _build_kill_rows(
        self, match_id: int, data: dict, has_replay_data: bool, unparseable: bool
    ) -> list[dict]:
        rampage_time, rampage_kill_times = self._detect_rampage(match_id, data)
        rampage_time_min = round(rampage_time / 60, 2) if rampage_time is not None else None
        rampage_kill_times_set = set(rampage_kill_times)
        rows: list[dict] = []

        for player in data.get("players") or []:
            killer_hero_id = player.get("hero_id")
            for kill in player.get("kills_log") or []:
                time_value = kill.get("time")
                victim_key = kill.get("key")
                if time_value is None or not victim_key:
                    continue

                victim_hero_id = self._victim_hero_id(victim_key)
                rows.append(
                    {
                        "match_id": match_id,
                        "kill_time": time_value,
                        "kill_time_min": round(time_value / 60, 2),
                        "killer_hero_id": killer_hero_id,
                        "killer_hero_name": self._hero_name(killer_hero_id),
                        "victim_hero_id": victim_hero_id,
                        "victim_hero_name": self._hero_name(victim_hero_id),
                        "rampage_time": rampage_time,
                        "rampage_time_min": rampage_time_min,
                        "is_rampage_kill": 1 if time_value in rampage_kill_times_set else 0,
                        "has_replay_data": 1 if has_replay_data else 0,
                        "unparseable": 1 if unparseable else 0,
                    }
                )

        return rows

    def _write_csv(self, csv_path: Path, fieldnames: list[str], rows: list[dict]) -> None:
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)

    def _write_excel(
        self,
        xlsx_path: Path,
        chat_rows: list[dict],
        kill_rows: list[dict],
        match_ids: list[int],
    ) -> None:
        wb = Workbook()
        chat_sheet = wb.active
        chat_sheet.title = "Chats"
        kill_sheet = wb.create_sheet("Kills")

        match_color_map = self._build_match_color_map(match_ids)

        self._style_sheet(chat_sheet, self.CHAT_FIELDNAMES, chat_rows, match_color_map)
        self._style_sheet(kill_sheet, self.KILL_FIELDNAMES, kill_rows, match_color_map)

        wb.save(xlsx_path)

    def _build_match_color_map(self, match_ids: list[int]) -> dict[int, int]:
        seen: list[int] = []
        for match_id in match_ids:
            if match_id not in seen:
                seen.append(match_id)
        return {match_id: idx % 2 for idx, match_id in enumerate(seen)}

    def _style_sheet(
        self,
        ws,
        fieldnames: list[str],
        rows: list[dict],
        match_color_map: dict[int, int] | None = None,
    ) -> None:
        header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        header_font = Font(color="F8FAFC", bold=True, size=11)
        border_color = Side(style="thin", color="475569")
        light_green_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
        darker_green_fill = PatternFill(start_color="A7F3D0", end_color="A7F3D0", fill_type="solid")
        fills = [light_green_fill, darker_green_fill]
        font = Font(color="1F2937", size=10)

        for col_idx, fieldname in enumerate(fieldnames, start=1):
            cell = ws.cell(row=1, column=col_idx, value=fieldname)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=border_color)

        for row_idx, row in enumerate(rows, start=2):
            match_id = row.get("match_id")
            color_idx = 0
            if match_color_map is not None and match_id in match_color_map:
                color_idx = match_color_map[match_id]
            fill = fills[color_idx]
            for col_idx, fieldname in enumerate(fieldnames, start=1):
                value = row.get(fieldname)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.fill = fill
                cell.font = font
                cell.alignment = Alignment(vertical="center")

        for col_idx, fieldname in enumerate(fieldnames, start=1):
            max_length = len(fieldname)
            for row in rows:
                value = str(row.get(fieldname) or "")
                if len(value) > max_length:
                    max_length = min(len(value), 60)
            ws.column_dimensions[get_column_letter(col_idx)].width = max_length + 2

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
